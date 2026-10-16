#!/usr/bin/env python3
"""Markdown → docx。覆盖报告里会用到的子集：标题、段落、粗体/斜体、无序/有序列表、表格、引用、代码块、图片。
用法：md_to_docx.py in.md out.docx"""
from __future__ import annotations

import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, Inches
except ImportError:
    print("missing dependency. run: pip install python-docx", file=sys.stderr)
    sys.exit(2)

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def add_inline(par, text: str) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`"):
            r = par.add_run(part[1:-1]); r.font.name = "Consolas"
        elif part.startswith("*"):
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def flush_table(doc, rows: list[list[str]]) -> None:
    rows = [r for r in rows if not all(re.fullmatch(r":?-+:?", c.strip()) for c in r)]
    if not rows:
        return
    t = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
    t.style = "Light Grid Accent 1"
    for i, r in enumerate(rows):
        for j, c in enumerate(r):
            cell = t.cell(i, j)
            cell.text = ""
            add_inline(cell.paragraphs[0], c.strip())
            if i == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True


def main() -> int:
    if len(sys.argv) != 3:
        print("usage: md_to_docx.py in.md out.docx", file=sys.stderr)
        return 2
    src, out = Path(sys.argv[1]), Path(sys.argv[2])
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10.5)
    table: list[list[str]] = []
    code: list[str] | None = None
    for raw in src.read_text(encoding="utf-8").splitlines():
        line = raw.rstrip()
        if code is not None:
            if line.startswith("```"):
                p = doc.add_paragraph(); r = p.add_run("\n".join(code)); r.font.name = "Consolas"; r.font.size = Pt(9)
                code = None
            else:
                code.append(raw)
            continue
        if line.startswith("```"):
            code = []; continue
        if line.startswith("|"):
            table.append([c for c in line.strip("|").split("|")]); continue
        if table:
            flush_table(doc, table); table = []
        if not line.strip():
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", line)
        if m:
            doc.add_heading(m.group(2).strip(), level=len(m.group(1))); continue
        m = re.match(r"^!\[[^\]]*\]\(([^)]+)\)", line)
        if m and (src.parent / m.group(1)).is_file():
            doc.add_picture(str(src.parent / m.group(1)), width=Inches(6)); continue
        if line.startswith(">"):
            p = doc.add_paragraph(style="Intense Quote"); add_inline(p, line.lstrip("> ")); continue
        m = re.match(r"^\s*[-*]\s+(.*)", line)
        if m:
            p = doc.add_paragraph(style="List Bullet"); add_inline(p, m.group(1)); continue
        m = re.match(r"^\s*\d+[.)]\s+(.*)", line)
        if m:
            p = doc.add_paragraph(style="List Number"); add_inline(p, m.group(1)); continue
        p = doc.add_paragraph(); add_inline(p, line)
    if table:
        flush_table(doc, table)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(f"wrote {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
