#!/usr/bin/env python3
"""JSON 表格 → xlsx。输入是一个 sheet 列表或单个 sheet：
  [{"name": "因子", "columns": ["因子", "数值", "窗口"], "rows": [["pe_ttm", 26.5, "TTM"], ...], "note": "数据来源：..."}]
用法：table_to_xlsx.py tables.json out.xlsx"""
from __future__ import annotations

import json
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    print("missing dependency. run: pip install openpyxl", file=sys.stderr)
    sys.exit(2)


def main() -> int:
    if len(sys.argv) != 3:
        print("usage: table_to_xlsx.py tables.json out.xlsx", file=sys.stderr)
        return 2
    spec = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    sheets = spec if isinstance(spec, list) else [spec]
    wb = Workbook()
    wb.remove(wb.active)
    for i, sh in enumerate(sheets):
        ws = wb.create_sheet((sh.get("name") or f"Sheet{i + 1}")[:31])
        cols = sh.get("columns", [])
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="E3E8FA")
            c.alignment = Alignment(horizontal="center")
        for row in sh.get("rows", []):
            ws.append(list(row))
        for j, col in enumerate(cols, start=1):
            width = max([len(str(col))] + [len(str(r[j - 1])) for r in sh.get("rows", []) if j - 1 < len(r)]) + 2
            ws.column_dimensions[get_column_letter(j)].width = min(max(width, 8), 50)
        ws.freeze_panes = "A2"
        if sh.get("note"):
            ws.append([])
            ws.append([sh["note"]])
            ws.cell(ws.max_row, 1).font = Font(italic=True, color="8A9099")
    out = Path(sys.argv[2])
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"wrote {out} ({len(sheets)} sheets)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
